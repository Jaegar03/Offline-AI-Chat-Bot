import os
import tempfile
from flask import Flask, request, send_file
from werkzeug.utils import secure_filename
from docx import Document
import PyPDF2
from fpdf import FPDF

app = Flask(__name__)

# Path to Unicode font (must exist)
FONT_PATH = os.path.join(os.path.dirname(__file__), "DejaVuSans.ttf")

def generate_ieee_pdf(text):
    pdf = FPDF()
    pdf.add_page()
    pdf.add_font("DejaVu", "", FONT_PATH, uni=True)
    pdf.set_font("DejaVu", size=12)

    # Header
    pdf.set_font("DejaVu", style='B', size=14)
    pdf.cell(200, 10, "IEEE Conference Paper", ln=True, align='C')
    pdf.ln(10)

    # Body text
    pdf.set_font("DejaVu", size=12)
    for line in text.split("\n"):
        pdf.multi_cell(0, 10, line)

    temp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    pdf.output(temp_pdf.name)
    return temp_pdf.name

@app.route("/upload", methods=["POST"])
def upload_file():
    try:
        uploaded_file = request.files.get("file")
        if not uploaded_file:
            return "No file uploaded", 400

        filename = secure_filename(uploaded_file.filename)

        # === Extract Text ===
        if filename.endswith(".txt"):
            text = uploaded_file.read().decode("utf-8")
        elif filename.endswith(".pdf"):
            reader = PyPDF2.PdfReader(uploaded_file)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif filename.endswith(".docx"):
            with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
                tmp.write(uploaded_file.read())
                tmp.flush()
                doc = Document(tmp.name)
            text = "\n".join(para.text for para in doc.paragraphs)
        else:
            return "Unsupported file type", 400

        # === Generate IEEE Styled PDF ===
        pdf_path = generate_ieee_pdf(text)
        return send_file(pdf_path, as_attachment=True, download_name="ieee_formatted.pdf")

    except Exception as e:
        print("❌ Flask error:", str(e))
        return f"Internal server error: {str(e)}", 500

if __name__ == "__main__":
    app.run(debug=True)
